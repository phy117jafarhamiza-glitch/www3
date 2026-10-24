# -*- coding: utf-8 -*-
"""
مولّد البرنامج التأهيلي الفردي (8 أسابيع × 3 وحدات) لعضلات الرقبة عند لاعبي الفنون القتالية.

وحدة بايثون نقية (دون streamlit). كل العتبات في DEFAULT_CFG «مبدئية قابلة للتعديل من الأخصائي»
والبرنامج الناتج مقترح ويخضع لموافقة الأخصائي/المشرف قبل التطبيق.

الواجهة:
    build_profile(subject, m, cfg=None) -> profile
    generate_program(profile, cfg=None) -> program
    coach_spec_for(ex, m, week, cfg=None) -> dict | None
    program_to_dataframe(program) / program_to_xlsx_bytes(program) / program_to_docx_bytes(program)

سلوك التباطؤ (ألم/خوف): factor < 1 => الأسبوع الفعّال effective_week = ceil(week*factor).
يُستعمل الأسبوع الفعّال فقط في اختيار المستوى (1..4) وجرعة المجموعات/التكرارات/الراحة وسقف المقاومة
(أي أن الحِمل الرئيسي مسقوف)، أما مدة الوحدات فتتبع الخطة الأصلية دائمًا وتُملأ بحجم أخف
(تمارين إضافية/مجموعات إضافية/راحة أطول) بدل رفع الشدة.

تقدير زمن التمرين (دقيقة):  sets * (reps * (hold_s أو tempo_s)) / 60  +  sets * rest_s / 60  +  0.5 (انتقال).
"""
import io
import math
from copy import deepcopy
from datetime import date

import pandas as pd

from utils.exercise_bank import BANK, BANK_BY_ID

DIRECTIONS = ["flex", "ext", "lat_r", "lat_l", "rot_r", "rot_l"]
DIRECTION_AR = {"flex": "ثني أمامي", "ext": "بسط خلفي", "lat_r": "ميل أيمن", "lat_l": "ميل أيسر",
                "rot_r": "دوران أيمن", "rot_l": "دوران أيسر"}

ALL_TAGS = (["mobility_" + d for d in DIRECTIONS] + ["deep_flexor"] + ["str_" + d for d in DIRECTIONS] +
            ["scapular", "core", "proprio", "reaction", "relax", "fear", "sport", "eccentric"])

DISCLAIMER = ("البرنامج مقترح آليًا بحسب نتائج الاختبارات، وكل العتبات مبدئية قابلة للتعديل من الأخصائي؛ "
              "ولا يُطبَّق إلا بعد مراجعة واعتماد الأخصائي/المشرف.")

# --------------------------------------------------------------------------------------
DEFAULT_CFG = {
    # --- كل القيم أدناه مبدئية قابلة للتعديل من الأخصائي ---
    "norm_rom": {"flex": 50.0, "ext": 60.0, "lat": 45.0, "rot": 80.0},   # درجات (مرجع مبدئي)
    "rom_deficit_frac": 0.20,       # عجز المدى >= 20% من المرجع
    "rom_severe_frac": 0.40,        # عجز شديد >= 40%
    "fe_ratio_ref": 0.65,           # نسبة قوة الثني/البسط المرجعية
    "fe_ratio_tol": 0.15,           # هامش التفاوت المسموح حول النسبة
    "asym_pct": 10.0,               # عدم تماثل يمين/يسار % (قوة أو مدى)
    "asym_strong_pct": 20.0,        # عدم تماثل كبير
    "weak_rel": 0.80,               # اتجاه ضعيف إن كان < 80% من متوسط اتجاهات الميل/الدوران
    "pain_severe": 6.0,             # NPRS >= 6 شديد
    "pain_moderate": 4.0,           # 4–5 متوسط
    "tsk_high": 26.0,               # TSK-11 مرتفع
    "kin_guard_high": 50.0,         # مؤشر الحذر الحركي (فيديو) مرتفع
    "emg_scm_scalene_pct": 20.0,    # متوسط SCM/Scalene >= 20 %MVC => نشاط سطحي زائد
    "emg_trap_pct": 30.0,           # الشبه منحرفة العلوية >= 30 %MVC
    "rt_slow_ms": 500.0,            # زمن استجابة بطيء
    "rt_acc_min": 85.0,             # الدقة < 85% => ضعيفة
    # --- التقدم ---
    "factor_severe_pain": 0.6, "factor_moderate_pain": 0.8, "factor_high_fear": -0.1, "factor_min": 0.4,
    "base_minutes": [15, 25, 35, 45, 55, 65, 75, 90],
    "sets": [2, 2, 3, 3, 3, 3, 4, 4],
    "reps": [8, 10, 12, 12, 12, 14, 15, 15],
    "rest_s": [60, 45, 45, 45, 30, 30, 30, 30],
    "resistance_cap_pct": [30, 30, 45, 45, 60, 60, 70, 70],     # % من MVC
    "external_resistance_from_week": 3,                          # لا مقاومة خارجية في الأسبوعين 1–2
    "level_fraction": {1: 0.50, 2: 0.65, 3: 0.80, 4: 0.90},      # نسبة المدى المستهدف للمدرب المباشر
    "warmup_frac": 0.15, "cooldown_frac": 0.10, "min_block_min": 3,
    "max_per_category": 2,
    "unit_min": 15, "unit_max": 90,
}


def merge_cfg(cfg=None):
    c = deepcopy(DEFAULT_CFG)
    if cfg:
        for k, v in cfg.items():
            if isinstance(v, dict) and isinstance(c.get(k), dict):
                c[k].update(v)
            else:
                c[k] = v
    return c


# --------------------------------------------------------------------------------------
# الملف التعريفي للاعب
def _val(m, k):
    v = (m or {}).get(k)
    try:
        if v is None:
            return None
        v = float(v)
        return None if math.isnan(v) else v
    except (TypeError, ValueError):
        return None


def _side_ar(d):
    return "يمين" if d.endswith("_r") else ("يسار" if d.endswith("_l") else None)


def build_profile(subject, m, cfg=None):
    """يبني ملف اللاعب من نتائج الاختبارات: نتائج (findings) + أوزان أولوية للوسوم + معامل التقدم + احتياطات."""
    c = merge_cfg(cfg)
    subject = dict(subject or {})
    m = dict(m or {})
    w = {t: 0.5 for t in ALL_TAGS}          # أساس: كل الوسوم لها أولوية دنيا
    w["sport"] = 1.0
    reasons = {t: [] for t in ALL_TAGS}     # وسم -> أسباب عربية
    findings = []
    extra_set_tags = set()
    weak_str_tags = []
    flags = dict(flexor_weak=False, extensor_weak=False, emg_comp=False, reaction_slow=False, fear_high=False,
                 guarding=False, pain_level=None, pain_class="none", rom_deficit_dirs=[], asym={})
    precautions = ["أوقف التمرين فورًا عند الدوخة أو التنميل أو الصداع الحاد أو الضعف المفاجئ وأبلغ الأخصائي.",
                   "يُمنع أي حمل محوري ثقيل على الرقبة وكذلك «جسر الرقبة» الكامل؛ ولا تُجرى تمارين الاحتكاك إلا بإشراف.",
                   DISCLAIMER]

    def add(tag, inc, reason=None):
        w[tag] = min(3.0, w[tag] + inc)
        if reason and reason not in reasons[tag]:
            reasons[tag].append(reason)

    def finding(key, title, severity, text, tags=()):
        findings.append(dict(key=key, title=title, severity=severity, text=text, tags=list(tags)))

    # ---- القوة
    st = {d: _val(m, "str_" + d) for d in DIRECTIONS}
    measured_st = {d: v for d, v in st.items() if v is not None}
    if not measured_st:
        finding("strength", "القوة الإيزومترية", "unmeasured", "القوة الإيزومترية: لم يُقَس — تُطبَّق أولويات افتراضية متوازنة.")
    else:
        # النسبة ثني/بسط
        if st["flex"] is not None and st["ext"] not in (None, 0):
            ratio = st["flex"] / st["ext"]
            lo, hi = c["fe_ratio_ref"] - c["fe_ratio_tol"], c["fe_ratio_ref"] + c["fe_ratio_tol"]
            if ratio < lo:
                flags["flexor_weak"] = True
                r = f"ضعف نسبي في المثنيات (نسبة ثني/بسط {ratio:.2f} أقل من المرجع {c['fe_ratio_ref']:.2f})"
                add("str_flex", 2.0, r); add("deep_flexor", 2.5, r); add("eccentric", 1.0, r)
                extra_set_tags.add("str_flex"); weak_str_tags.append("str_flex")
                finding("fe_ratio", "توازن المثنيات والباسطات", "moderate",
                        f"نسبة قوة الثني إلى البسط {ratio:.2f} (المرجع المبدئي {c['fe_ratio_ref']:.2f} ± {c['fe_ratio_tol']:.2f}): "
                        "المثنيات أضعف نسبيًا، لذا تُعطى أولوية للمثنيات العميقة وتقوية الثني.", ["str_flex", "deep_flexor"])
            elif ratio > hi:
                flags["extensor_weak"] = True
                r = f"ضعف نسبي في الباسطات (نسبة ثني/بسط {ratio:.2f} أعلى من المرجع)"
                add("str_ext", 2.0, r); add("scapular", 1.0, r); add("core", 0.5, r)
                extra_set_tags.add("str_ext"); weak_str_tags.append("str_ext")
                finding("fe_ratio", "توازن المثنيات والباسطات", "moderate",
                        f"نسبة قوة الثني إلى البسط {ratio:.2f} أعلى من المرجع المبدئي: الباسطات أضعف نسبيًا، "
                        "فتُعطى أولوية لتقوية الباسطات وحزام الكتف.", ["str_ext", "scapular"])
            else:
                finding("fe_ratio", "توازن المثنيات والباسطات", "none",
                        f"نسبة الثني/البسط {ratio:.2f} ضمن الهامش المرجعي المبدئي.")
        else:
            finding("fe_ratio", "توازن المثنيات والباسطات", "unmeasured", "نسبة الثني/البسط: لم يُقَس أحد الاتجاهين.")
        # عدم تماثل يمين/يسار
        for base, nm in (("lat", "الميل الجانبي"), ("rot", "الدوران")):
            r_, l_ = st[base + "_r"], st[base + "_l"]
            if r_ is not None and l_ is not None and max(r_, l_) > 0:
                asym = abs(r_ - l_) / max(r_, l_) * 100
                flags["asym"]["str_" + base] = round(asym, 1)
                if asym >= c["asym_pct"]:
                    weak = base + ("_r" if r_ < l_ else "_l")
                    tag = "str_" + weak
                    inc = 2.5 if asym >= c["asym_strong_pct"] else 2.0
                    rs = f"ضعف {nm} جهة {'اليمين' if weak.endswith('_r') else 'اليسار'} (عدم تماثل {asym:.0f}%)"
                    add(tag, inc, rs); extra_set_tags.add(tag); weak_str_tags.append(tag)
                    finding("asym_str_" + base, f"تماثل قوة {nm}", "severe" if asym >= c["asym_strong_pct"] else "moderate",
                            f"فرق القوة بين اليمين واليسار في {nm} {asym:.0f}% (العتبة المبدئية {c['asym_pct']:.0f}%): "
                            f"تُمنح الجهة الأضعف ({'اليمين' if weak.endswith('_r') else 'اليسار'}) حجمًا إضافيًا (مجموعة زائدة).", [tag])
                else:
                    finding("asym_str_" + base, f"تماثل قوة {nm}", "none", f"فرق قوة {nm} يمين/يسار {asym:.0f}% ضمن الحد المبدئي.")
        # ضعف نسبي لاتجاه بالنسبة لمتوسط الاتجاهات الجانبية/الدورانية
        lr = {d: st[d] for d in ("lat_r", "lat_l", "rot_r", "rot_l") if st[d] is not None}
        if len(lr) >= 3:
            mean = sum(lr.values()) / len(lr)
            for d, v in lr.items():
                tag = "str_" + d
                if v < c["weak_rel"] * mean and tag not in weak_str_tags:
                    rs = f"قوة {DIRECTION_AR[d]} أقل من {c['weak_rel']*100:.0f}% من متوسط اتجاهات الميل والدوران"
                    add(tag, 1.5, rs); extra_set_tags.add(tag); weak_str_tags.append(tag)
                    finding("weak_" + d, f"قوة {DIRECTION_AR[d]}", "mild", rs + ".", [tag])

    # ---- مدى الحركة
    rom = {d: _val(m, "rom_" + d) for d in DIRECTIONS}
    if all(v is None for v in rom.values()):
        finding("rom", "مدى الحركة", "unmeasured", "مدى الحركة: لم يُقَس — يُستخدم إحماء حركي متوازن لكل الاتجاهات.")
    else:
        for d in DIRECTIONS:
            v = rom[d]
            if v is None:
                continue
            norm = c["norm_rom"][d.split("_")[0]]
            deficit = 1 - v / norm
            if deficit >= c["rom_deficit_frac"]:
                flags["rom_deficit_dirs"].append(d)
                sev = "severe" if deficit >= c["rom_severe_frac"] else "moderate"
                inc = 3.0 if sev == "severe" else 2.0
                rs = f"عجز في مدى {DIRECTION_AR[d]} ({v:.0f}° مقابل مرجع {norm:.0f}°، عجز {deficit*100:.0f}%)"
                add("mobility_" + d, inc, rs); extra_set_tags.add("mobility_" + d)
                finding("rom_" + d, f"مدى {DIRECTION_AR[d]}", sev, rs + ": حركية إضافية في هذا الاتجاه.", ["mobility_" + d])
        for base, nm in (("lat", "الميل الجانبي"), ("rot", "الدوران")):
            r_, l_ = rom[base + "_r"], rom[base + "_l"]
            if r_ is not None and l_ is not None and max(r_, l_) > 0:
                asym = abs(r_ - l_) / max(r_, l_) * 100
                flags["asym"]["rom_" + base] = round(asym, 1)
                if asym >= c["asym_pct"]:
                    lesser = base + ("_r" if r_ < l_ else "_l")
                    tag = "mobility_" + lesser
                    rs = f"عدم تماثل مدى {nm} ({asym:.0f}%) والجهة المحدودة {DIRECTION_AR[lesser]}"
                    add(tag, 2.0, rs); extra_set_tags.add(tag)
                    if lesser not in flags["rom_deficit_dirs"]:
                        flags["rom_deficit_dirs"].append(lesser)
                    finding("asym_rom_" + base, f"تماثل مدى {nm}", "moderate", rs + ".", [tag])
        if not flags["rom_deficit_dirs"]:
            finding("rom_ok", "مدى الحركة", "none", "لا عجز واضح في مدى الحركة بحسب العتبات المبدئية.")

    # ---- الألم
    pm, pr = _val(m, "nprs_move"), _val(m, "nprs_rest")
    pains = [p for p in (pm, pr) if p is not None]
    factor = 1.0
    if not pains:
        finding("pain", "الألم", "unmeasured", "الألم (NPRS): لم يُقَس — يُعامل بحذر كأن الألم غير معروف.")
    else:
        p = max(pains)
        flags["pain_level"] = p
        if p >= c["pain_severe"]:
            flags["pain_class"] = "severe"; factor = c["factor_severe_pain"]
            add("relax", 2.5, "ألم شديد"); add("fear", 0.5)
            precautions.append("ألم شديد: يُبطَّأ التقدم ويُسقَّف الحِمل، ويُوقف التمرين إن ارتفع الألم نقطتين عن بداية الجلسة.")
            finding("pain", "الألم", "severe", f"أعلى ألم مسجل {p:.0f}/10 (شديد ≥ {c['pain_severe']:.0f}): تقدم أبطأ ومقاومة منخفضة وتركيز على التهدئة.", ["relax"])
        elif p >= c["pain_moderate"]:
            flags["pain_class"] = "moderate"; factor = c["factor_moderate_pain"]
            add("relax", 1.5, "ألم متوسط")
            precautions.append("ألم متوسط: تقدم أبطأ نسبيًا ومراقبة الألم قبل وبعد كل وحدة.")
            finding("pain", "الألم", "moderate", f"أعلى ألم مسجل {p:.0f}/10 (متوسط): تقدم أبطأ قليلًا وإضافة تمارين تهدئة.", ["relax"])
        elif p > 0:
            flags["pain_class"] = "mild"
            finding("pain", "الألم", "mild", f"ألم خفيف {p:.0f}/10: تقدم اعتيادي مع المراقبة.")
        else:
            finding("pain", "الألم", "none", "لا ألم مسجل.")

    # ---- الخوف من الحركة
    tsk = _val(m, "tsk11")
    kin = _val(m, "kin_guard")
    if tsk is None:
        finding("tsk", "الخوف من الحركة (TSK-11)", "unmeasured", "TSK-11: لم يُقَس.")
    elif tsk >= c["tsk_high"]:
        flags["fear_high"] = True
        rs = f"خوف حركي مرتفع (TSK-11 = {tsk:.0f})"
        add("fear", 3.0, rs); add("relax", 1.0, rs)
        for d in DIRECTIONS:
            add("mobility_" + d, 0.5)
        precautions.append("خوف حركي مرتفع: يبدأ كل أسبوع بتعرّض تدريجي وتثقيف، وتُسجَّل درجة الخوف قبل وبعد.")
        finding("tsk", "الخوف من الحركة (TSK-11)", "severe" if tsk >= c["tsk_high"] + 8 else "moderate",
                f"{tsk:.0f} من 44 (العتبة المبدئية {c['tsk_high']:.0f}): تعرّض حركي متدرج وتثقيف الألم في كل وحدة.", ["fear"])
    else:
        finding("tsk", "الخوف من الحركة (TSK-11)", "none", f"{tsk:.0f} من 44 دون العتبة المبدئية.")
    if kin is not None:
        if kin >= c["kin_guard_high"]:
            flags["guarding"] = True
            rs = f"حذر حركي مرتفع في الفيديو (مؤشر {kin:.0f}/100)"
            add("fear", 1.5, rs); add("relax", 1.0, rs); add("proprio", 1.0, rs)
            finding("kin_guard", "الحذر الحركي (فيديو)", "moderate", rs + ": تُضاف تمارين تعرّض وتحكم هادئ.", ["fear", "proprio"])
        else:
            finding("kin_guard", "الحذر الحركي (فيديو)", "none", f"مؤشر الحذر {kin:.0f}/100 دون العتبة.")
    if flags["fear_high"] or flags["guarding"]:
        flags["fear_flag"] = True
        factor += c["factor_high_fear"]
    factor = round(max(c["factor_min"], min(1.0, factor)), 3)

    # ---- EMG
    es, ea, et = _val(m, "emg_scm"), _val(m, "emg_scalene"), _val(m, "emg_trap")
    if es is None and ea is None and et is None:
        finding("emg", "EMG أثناء CCFT", "unmeasured", "EMG: لم يُقَس.")
    else:
        sup = [v for v in (es, ea) if v is not None]
        comp = (sup and sum(sup) / len(sup) >= c["emg_scm_scalene_pct"]) or (et is not None and et >= c["emg_trap_pct"])
        if comp:
            flags["emg_comp"] = True
            rs = "نشاط عضلي سطحي زائد أثناء اختبار الثني القحفي الرقبي (تعويض)"
            add("deep_flexor", 3.0, rs); add("relax", 1.5, rs); add("scapular", 1.5, rs); add("proprio", 0.5, rs)
            precautions.append("تعويض عضلي سطحي: الجودة قبل الحِمل؛ تُتجنب المجهودات القصوى وتُراقب الشقرافية/القصية أثناء التنشيط.")
            finding("emg", "EMG أثناء CCFT", "moderate",
                    "SCM/Scalene: " + ("-" if not sup else f"{sum(sup)/len(sup):.0f}%") + " ، الشبه منحرفة: " +
                    ("-" if et is None else f"{et:.0f}%") + f" MVC (العتبات المبدئية {c['emg_scm_scalene_pct']:.0f}% / {c['emg_trap_pct']:.0f}%): "
                    "تُعطى أولوية لتنشيط المثنيات العميقة والاسترخاء.", ["deep_flexor", "relax"])
        else:
            finding("emg", "EMG أثناء CCFT", "none", "لا نشاط سطحي زائد بحسب العتبات المبدئية.")

    # ---- زمن الاستجابة
    rt, acc = _val(m, "rt_ms"), _val(m, "rt_acc")
    if rt is None and acc is None:
        finding("reaction", "زمن الاستجابة", "unmeasured", "زمن الاستجابة/دقته: لم يُقَس.")
    else:
        slow = rt is not None and rt >= c["rt_slow_ms"]
        inacc = acc is not None and acc < c["rt_acc_min"]
        if slow or inacc:
            flags["reaction_slow"] = True
            parts = []
            if slow:
                parts.append(f"زمن الاستجابة {rt:.0f} مللي ث (العتبة {c['rt_slow_ms']:.0f})")
            if inacc:
                parts.append(f"الدقة {acc:.0f}% (الحد الأدنى {c['rt_acc_min']:.0f}%)")
            rs = " و".join(parts)
            add("reaction", 3.0, rs); add("proprio", 2.0, rs)
            finding("reaction", "زمن الاستجابة", "moderate", rs + ": تُدرج تمارين استجابة وتتبع وحس عميق.", ["reaction", "proprio"])
        else:
            finding("reaction", "زمن الاستجابة", "none", "زمن الاستجابة ودقته ضمن الحدود المبدئية.")

    # ---- الرياضة
    sport = str(subject.get("sport") or "")
    if sport:
        add("sport", 0.5, f"رياضة اللاعب: {sport}")

    n_tests = sum(1 for k in ["str_" + d for d in DIRECTIONS] + ["rom_" + d for d in DIRECTIONS] +
                  ["nprs_move", "nprs_rest", "tsk11", "kin_guard", "emg_scm", "emg_scalene", "emg_trap", "rt_ms", "rt_acc"]
                  if _val(m, k) is not None)
    return dict(subject=subject, measurements={k: v for k, v in m.items()}, findings=findings, weights=w,
                tag_reasons=reasons, progression_factor=factor, precautions=precautions, flags=flags,
                extra_set_tags=sorted(extra_set_tags), weak_str_tags=weak_str_tags, sport=sport, n_tests=n_tests,
                notes=["لم يُقَس: " + f["title"] for f in findings if f["severity"] == "unmeasured"])


# --------------------------------------------------------------------------------------
# مواصفات مدرب الكاميرا
def _level_of(week):
    return max(1, min(4, math.ceil(week / 2)))


def coach_spec_for(ex, m, week, cfg=None):
    """يحوّل مواصفات cam للتمرين إلى مواصفات ExerciseCoach (target/safe بالدرجات). None إن لم يكن قابلًا للتتبع."""
    c = merge_cfg(cfg)
    if isinstance(ex, int):
        ex = BANK_BY_ID[ex]
    cam = ex.get("cam")
    if not cam:
        return None
    level = _level_of(week)
    frac = c["level_fraction"][level]
    rom_key = cam.get("rom_key")
    rom = _val(m, rom_key) if rom_key else None
    if cam["mode"] == "neutral_hold":
        # هنا target = التفاوت المسموح عن الوضع المحايد (أضيق مع المستوى)
        tol = {1: 1.25, 2: 1.1, 3: 1.0, 4: 0.9}[level]
        target = max(5.0, cam["default_target"] * tol)
        safe = target + cam["safe_extra"]
    elif rom is not None and rom > 0:
        target = max(5.0, min(frac * rom, 0.95 * rom))
        safe = max(target, min(target + cam["safe_extra"], rom + 5.0))
    else:
        target = max(5.0, cam["default_target"] * frac)
        safe = target + cam["safe_extra"]
    spec = dict(mode=cam["mode"], axis=cam["axis"], sign=cam["sign"], view=cam["view"],
                min_rep_s=cam["min_rep_s"], target=round(target, 1), safe=round(safe, 1), level=level,
                fraction=frac, rom_key=rom_key, rom_used=rom)
    if cam["mode"] in ("hold", "neutral_hold"):
        spec["hold_s"] = cam.get("hold_s", ex.get("hold_s") or 20.0)
    return spec


# --------------------------------------------------------------------------------------
# بناء البرنامج
UNIT_NAMES = {1: "تقوية وتحمل", 2: "تحكم حركي وحس عميق وتفاعل", 3: "مختلطة ونوعية رياضية"}
_UNIT_CAT_BONUS = {
    1: {"قوة إيزومترية": 1.0, "مقاومة ديناميكية": 1.0, "باسطات الرقبة": 1.0, "حزام الكتف": 1.0},
    2: {"تنشيط المثنيات العميقة": 1.0, "الحس العميق": 1.0, "تحكم عصبي عضلي واستجابة": 1.2, "ثبات الجذع": 1.0},
    3: {"نوعي رياضي": 1.5, "حركية المدى": 0.5, "تعرّض تدريجي وتثقيف": 0.5, "الحس العميق": 0.6, "ثبات الجذع": 0.6},
}
_SPORT_ROOTS = {"مصارع": ["مصارع"], "جودو": ["جودو", "جيدو", "judo"], "ملاكم": ["ملاكم", "boxing"],
                "تايكوندو": ["تايكو", "taekwondo"]}


def _round5(x):
    return int(math.floor(x / 5.0 + 0.5) * 5)


def unit_minutes_for_week(week, cfg=None):
    c = merge_cfg(cfg)
    base = c["base_minutes"][week - 1]
    clamp = lambda x: max(c["unit_min"], min(c["unit_max"], x))
    return [clamp(base), clamp(_round5(0.9 * base)), clamp(base)]


def _sport_ok(ex, sport):
    if not ex.get("sports"):
        return True
    s = (sport or "").lower()
    return any(r in s for r in ex["sports"])


def _per_rep(ex_or_e):
    return ex_or_e.get("hold_s") or ex_or_e.get("tempo_s") or 3.0


def _emin(e):
    return e["sets"] * e["reps"] * _per_rep(e) / 60.0 + e["sets"] * e["rest_s"] / 60.0 + 0.5


def _score(ex, prof, unit_idx, level, prev_pen):
    s = sum(prof["weights"].get(t, 0.0) for t in ex["tags"])
    s += _UNIT_CAT_BONUS[unit_idx].get(ex["category"], 0.0)
    if unit_idx == 3:
        s += 0.3
    s += 0.1 * (level - min(ex["levels"]))
    s -= prev_pen.get(ex["id"], 0.0)
    return s


def _rationale(ex, prof):
    # أعلى الوسوم وزنًا التي لها سبب مرتبط بنتائج اللاعب
    cand = sorted([t for t in ex["tags"] if prof["tag_reasons"].get(t)], key=lambda t: -prof["weights"][t])
    if cand:
        rs = []
        for t in cand[:2]:
            for r in prof["tag_reasons"][t][:1]:
                if r not in rs:
                    rs.append(r)
        return "مختار بسبب: " + "؛ ".join(rs) + "."
    if "sport" in ex["tags"]:
        return "نوعي لرياضة اللاعب" + (f" ({prof['sport']})" if prof.get("sport") else "") + " بوضعيات آمنة وبلا جسر رقبة."
    return "جزء من البنية الأساسية المتوازنة للبرنامج (" + ex["category"] + ")."


def _make_entries(ex, prof, m, ctx, block, cfg):
    """ينشئ 1 أو 2 من مدخلات التمرين (لجهتين للتمارين الأحادية العامة)."""
    sets, reps = ctx["sets"], ctx["reps"]
    hold = ex.get("hold_s")
    rest = ctx["rest"]
    if ex.get("dose"):
        sets, reps = ex["dose"]["sets"], ex["dose"]["reps"]
        hold = ex["dose"].get("hold_s", hold)
        rest = 15 if ex["category"] == "استرخاء وإطالة" else 10
    else:
        if ex.get("rep_scale"):
            reps = max(3, int(round(reps * ex["rep_scale"])))
        if block == "warmup":
            sets = min(sets, 2)
    side_tag = next((t for t in ex["tags"] if t.endswith(("_r", "_l")) and t.split("_")[0] in ("mobility", "str")), None)
    if side_tag:
        sides = [_side_ar(side_tag)]
    elif ex.get("unilateral"):
        sides = ["يمين", "يسار"]
    else:
        sides = [None]
    extra = 1 if (set(ex["tags"]) & set(prof["extra_set_tags"]) and not ex.get("dose")) else 0
    pct = 0
    if ex["resisted"]:
        pct = ctx["cap"]
        if ex.get("max_pct") is not None:
            pct = min(pct, ex["max_pct"])
    out = []
    for sd in sides:
        out.append(dict(id=ex["id"], name=ex["name"], category=ex["category"], sets=int(sets + extra), reps=int(reps),
                        hold_s=hold, rest_s=int(rest), resistance_pct_mvc=int(pct), minutes=0.0, side=sd,
                        rationale=_rationale(ex, prof), cam=coach_spec_for(ex, m, ctx["eff_week"], cfg),
                        instructions=list(ex["instructions"]), common_errors=list(ex["common_errors"]),
                        equipment=ex["equipment"], tempo_s=ex["tempo_s"], phase=block,
                        _max_sets=int(sets + extra + 2), _base_sets=int(sets + extra)))
    for e in out:
        e["minutes"] = _emin(e)
    return out


def _item_min(item):
    return sum(_emin(e) for e in item["entries"])


def _fill(items, forced, budget, cap_cat):
    """يملأ كتلة زمنية: إلزامي ثم جشع بالدرجة مع حد فئات ثم مجموعات إضافية."""
    chosen, used, cats = [], 0.0, {}

    def can_add(it, tol=0.4):
        t = _item_min(it)
        if used + t <= budget + tol:
            return True
        # تقليص المجموعات حتى 1
        saved = [e["sets"] for e in it["entries"]]
        while any(e["sets"] > 1 for e in it["entries"]) and used + _item_min(it) > budget + tol:
            for e in it["entries"]:
                e["sets"] = max(1, e["sets"] - 1)
        ok = used + _item_min(it) <= budget + tol
        if not ok:
            for e, s in zip(it["entries"], saved):
                e["sets"] = s
        return ok

    def add(it):
        nonlocal used
        chosen.append(it)
        used += _item_min(it)
        cats[it["cat"]] = cats.get(it["cat"], 0) + 1

    for it in forced:
        if it in chosen:
            continue
        if can_add(it, 0.6) or not chosen:
            if not chosen and used + _item_min(it) > budget + 0.6:
                for e in it["entries"]:
                    e["sets"] = 1
            add(it)
    for extra_cap in (0, 1, 2, 4):
        if budget - used < 1.0:
            break
        for it in items:
            if it in chosen or cats.get(it["cat"], 0) >= cap_cat + extra_cap:
                continue
            if budget - used < 1.0:
                break
            if can_add(it):
                add(it)
    # مجموعات إضافية بحسب الأولوية
    changed = True
    while changed and budget - used > 0.3:
        changed = False
        for it in sorted(chosen, key=lambda x: (-x["score"], x["id"])):
            if any(e["sets"] >= e["_max_sets"] for e in it["entries"]):
                continue
            inc = sum(e["reps"] * _per_rep(e) / 60.0 + e["rest_s"] / 60.0 for e in it["entries"])
            if inc <= budget - used + 0.4:
                for e in it["entries"]:
                    e["sets"] += 1
                used += inc
                changed = True
    return chosen


def _adjust_rest(entries, delta_min):
    """يمتص فرق الزمن المتبقي بتعديل أزمنة الراحة (10..120 ث) دون تغيير الشدة."""
    if abs(delta_min) < 0.15 or not entries:
        return
    tot_sets = sum(e["sets"] for e in entries)
    per_set = delta_min * 60.0 / tot_sets
    for e in entries:
        e["rest_s"] = int(max(10, min(120, round(e["rest_s"] + per_set))))


def generate_program(profile, cfg=None):
    c = merge_cfg(cfg)
    prof = profile
    m = prof.get("measurements", {})
    factor = prof["progression_factor"]
    sport = prof.get("sport", "")
    weeks = []
    prev_week_ids = {1: set(), 2: set(), 3: set()}
    for week in range(1, 9):
        eff = min(8, max(1, math.ceil(round(week * factor, 6))))
        level = _level_of(eff)
        ctx = dict(sets=c["sets"][eff - 1], reps=c["reps"][eff - 1], rest=c["rest_s"][eff - 1],
                   cap=c["resistance_cap_pct"][eff - 1], eff_week=eff)
        ext_ok = week >= c["external_resistance_from_week"]
        eligible = [e for e in BANK if level in e["levels"] and _sport_ok(e, sport) and (ext_ok or not e["external"])]
        umins = unit_minutes_for_week(week, c)
        units, week_prev = [], {}
        this_ids = {1: set(), 2: set(), 3: set()}
        for ui in (1, 2, 3):
            T = umins[ui - 1]
            wu = max(c["min_block_min"], int(round(c["warmup_frac"] * T)))
            cu = max(c["min_block_min"], int(round(c["cooldown_frac"] * T)))
            mu = T - wu - cu
            pen = {}
            for pj in range(1, ui):
                for i in this_ids[pj]:
                    pen[i] = pen.get(i, 0) + (2.0 if pj == ui - 1 else 1.5)
            for i in prev_week_ids[ui]:
                pen[i] = pen.get(i, 0) + 0.4

            def mk_item(ex, blk):
                ents = _make_entries(ex, prof, m, ctx, blk, c)
                return dict(id=ex["id"], ex=ex, cat=ex["category"], entries=ents,
                            score=_score(ex, prof, ui, level, pen))

            def ranked(blk):
                its = [mk_item(e, blk) for e in eligible if e["phase"] == blk]
                return sorted(its, key=lambda x: (-x["score"], x["id"]))

            # ---------------- الجزء الرئيسي
            main_all = ranked("main")
            fl = prof["flags"]
            forced = []

            def best(pred):
                for it in main_all:
                    if pred(it["ex"]) and it not in forced:
                        return it
                return None
            if fl.get("fear_high") or fl.get("guarding"):
                forced.append(best(lambda e: "fear" in e["tags"]))
            if fl["emg_comp"] or fl["flexor_weak"]:
                forced.append(best(lambda e: "deep_flexor" in e["tags"]))
            if fl["reaction_slow"]:
                forced.append(best(lambda e: "reaction" in e["tags"] and e["category"] != "نوعي رياضي"))
                forced.append(best(lambda e: "proprio" in e["tags"] and e["category"] != "نوعي رياضي"))
            wst = prof["weak_str_tags"]
            for k in range(min(2, len(wst))):
                tg = wst[(ui - 1 + k) % len(wst)]
                forced.append(best(lambda e, tg=tg: tg in e["tags"] and e["resisted"]))
            forced = [f for f in forced if f]
            main_items = _fill(main_all, forced, mu, c["max_per_category"])
            # ---------------- الإحماء
            warm_all = ranked("warmup")
            wforced = []
            for d in prof["flags"]["rom_deficit_dirs"]:
                it = next((x for x in warm_all if "mobility_" + d in x["ex"]["tags"]), None)
                if it:
                    wforced.append(it)
            if wforced:  # تدوير الاتجاهات الناقصة بين الوحدات حتى تُغطّى كلها خلال الأسبوع
                off = (ui - 1) % len(wforced)
                wforced = wforced[off:] + wforced[:off]
            mob = next((x for x in warm_all if x["cat"] == "حركية المدى" and x not in wforced), None)
            if not wforced and mob:
                wforced.append(mob)
            warm_items = _fill(warm_all, wforced, wu, 6)
            # ---------------- التهدئة
            cool_all = ranked("cooldown")
            cforced = [x for x in cool_all if x["id"] == 36]
            if fl.get("fear_high") or fl.get("guarding"):
                cforced += [x for x in cool_all if x["id"] == 42]
            st = next((x for x in cool_all if x["cat"] == "استرخاء وإطالة" and x["id"] != 36 and x not in cforced), None)
            if st:
                cforced.append(st)
                if ui == 3:  # الوحدة 3: الإطالة أولًا (تنويع الوحدات عند ضيق الوقت)
                    cforced = [st] + [x for x in cforced if x is not st]
            cool_items = _fill(cool_all, cforced, cu, 4)

            def flat(items):
                out = []
                for it in items:
                    out.extend(it["entries"])
                return out
            W, M, C = flat(warm_items), flat(main_items), flat(cool_items)
            # امتصاص الفرق على أزمنة الراحة لكل كتلة ثم على الوحدة
            for blk, bud in ((W, wu), (M, mu), (C, cu)):
                _adjust_rest(blk, bud - sum(_emin(e) for e in blk))
            diff = T - sum(_emin(e) for e in W + M + C)
            _adjust_rest(M, diff)
            for e in W + M + C:
                e["minutes"] = round(_emin(e), 2)

            def clean(es):
                return [{k: v for k, v in e.items() if not k.startswith("_")} for e in es]
            ids = {e["id"] for e in W + M + C}
            this_ids[ui] = ids
            units.append(dict(unit=ui, name=UNIT_NAMES[ui], focus=UNIT_NAMES[ui], minutes=T,
                              est_minutes=round(sum(e["minutes"] for e in W + M + C), 1),
                              block_minutes=dict(warmup=wu, main=mu, cooldown=cu),
                              warmup=clean(W), main=clean(M), cooldown=clean(C)))
        prev_week_ids = this_ids
        weeks.append(dict(week=week, level=level, effective_week=eff, unit_minutes=umins, units=units,
                          dose=dict(sets=ctx["sets"], reps=ctx["reps"], rest_s=ctx["rest"],
                                    resistance_cap_pct=ctx["cap"]),
                          external_resistance_allowed=ext_ok))
    notes = [DISCLAIMER,
             "الوحدات الثلاث أسبوعيًا بمدة 15 إلى 90 دقيقة وتتدرج من السهل إلى الصعب: 15، 25، 35، 45، 55، 65، 75، 90 دقيقة للوحدتين 1 و3.",
             "لا مقاومة خارجية (مطاط/شريك/جهاز) في الأسبوعين 1–2؛ والمقاومة تُضبط كنسبة من MVC الأساسية ولا تتجاوز سقف الأسبوع.",
             "عند ألم شديد أو خوف حركي مرتفع يُبطَّأ التقدم (معامل %.2f): تبقى مدة الوحدة كما هي لكن بشدة أخف وحجم/راحة أكبر." % factor]
    for p in prof["precautions"][:-1]:
        if p not in notes:
            notes.append(p)
    meta = dict(name=prof["subject"].get("name"), code=prof["subject"].get("code"),
                group=prof["subject"].get("group_name"), sport=prof.get("sport"),
                created=date.today().isoformat(), progression_factor=factor, n_tests=prof["n_tests"],
                total_units=24, weeks=8, units_per_week=3, version="1.0", disclaimer=DISCLAIMER)
    return dict(meta=meta, profile=prof, weeks=weeks, notes=notes)


# --------------------------------------------------------------------------------------
# التصدير
_COLS = ["الأسبوع", "المستوى", "الوحدة", "اسم الوحدة", "مدة الوحدة (د)", "الجزء", "التمرين", "الجانب",
         "المجموعات", "التكرارات", "الثبات (ث)", "الراحة (ث)", "المقاومة (% من MVC)", "المدة (د)",
         "المعدات", "سبب الاختيار", "طريقة الأداء", "الأخطاء الشائعة"]
_BLK_AR = {"warmup": "إحماء", "main": "رئيسي", "cooldown": "تهدئة"}


def program_to_dataframe(program):
    rows = []
    for wk in program["weeks"]:
        for u in wk["units"]:
            for blk in ("warmup", "main", "cooldown"):
                for e in u[blk]:
                    rows.append([wk["week"], wk["level"], u["unit"], u["name"], u["minutes"], _BLK_AR[blk], e["name"],
                                 e["side"] or "—", e["sets"], e["reps"], e["hold_s"] if e["hold_s"] else "—",
                                 e["rest_s"], e["resistance_pct_mvc"] if e["resistance_pct_mvc"] else "—", e["minutes"],
                                 e["equipment"], e["rationale"], " | ".join(f"{i+1}) {s}" for i, s in enumerate(e["instructions"])),
                                 " | ".join(e["common_errors"])])
    return pd.DataFrame(rows, columns=_COLS)


_SEV_AR = {"none": "طبيعي", "mild": "خفيف", "moderate": "متوسط", "severe": "شديد", "unmeasured": "لم يُقَس"}


def program_to_xlsx_bytes(program):
    df = program_to_dataframe(program)
    meta, prof = program["meta"], program["profile"]
    summ = [["الاسم", meta.get("name")], ["الرمز", meta.get("code")], ["المجموعة", meta.get("group")],
            ["الرياضة", meta.get("sport")], ["تاريخ الإنشاء", meta.get("created")],
            ["معامل التقدم", meta.get("progression_factor")], ["عدد الاختبارات المدخلة", meta.get("n_tests")],
            ["ملاحظة", meta.get("disclaimer")], ["", ""], ["نتائج الاختبارات والقواعد", "الدرجة / التفسير"]]
    for f in prof["findings"]:
        summ.append([f["title"] + " [" + _SEV_AR[f["severity"]] + "]", f["text"]])
    summ.append(["", ""])
    summ.append(["الاحتياطات", ""])
    for p in prof["precautions"]:
        summ.append(["•", p])
    summ.append(["", ""])
    summ.append(["الأسبوع", "المستوى / دقائق الوحدات / الجرعة"])
    for wk in program["weeks"]:
        d = wk["dose"]
        summ.append([wk["week"], f"المستوى {wk['level']} | الدقائق {wk['unit_minutes']} | مجموعات {d['sets']} × تكرارات {d['reps']} | "
                                 f"راحة {d['rest_s']} ث | سقف المقاومة {d['resistance_cap_pct']}% MVC"])
    sdf = pd.DataFrame(summ, columns=["البند", "القيمة"])
    bio = io.BytesIO()
    with pd.ExcelWriter(bio, engine="openpyxl") as xw:
        sdf.to_excel(xw, sheet_name="الملخص", index=False)
        df.to_excel(xw, sheet_name="البرنامج الكامل", index=False)
        for wk in program["weeks"]:
            df[df["الأسبوع"] == wk["week"]].to_excel(xw, sheet_name=f"الأسبوع {wk['week']}", index=False)
        for ws in xw.book.worksheets:
            ws.sheet_view.rightToLeft = True
            for col in ws.columns:
                width = max((len(str(cell.value)) if cell.value is not None else 0) for cell in col[:60])
                ws.column_dimensions[col[0].column_letter].width = min(60, max(10, width * 0.9))
    return bio.getvalue()


def _rtl_par(par):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    pPr = par._p.get_or_add_pPr()
    if pPr.find(qn("w:bidi")) is None:
        pPr.append(OxmlElement("w:bidi"))
    for r in par.runs:
        rPr = r._r.get_or_add_rPr()
        if rPr.find(qn("w:rtl")) is None:
            rPr.append(OxmlElement("w:rtl"))
    return par


def program_to_docx_bytes(program):
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt, Cm
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(1.5)
    sec.top_margin = sec.bottom_margin = Cm(1.5)
    st = doc.styles["Normal"]
    st.font.name = "Arial"
    st.font.size = Pt(10)
    rpr = st.element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.append(rf)
    for a in ("w:ascii", "w:hAnsi", "w:cs"):
        rf.set(qn(a), "Arial")

    def para(text, bold=False, size=None):
        p = doc.add_paragraph()
        r = p.add_run(text)
        r.bold = bold
        if size:
            r.font.size = Pt(size)
        return _rtl_par(p)

    def head(text, lvl):
        h = doc.add_heading(text, level=lvl)
        for r in h.runs:
            rf2 = r._r.get_or_add_rPr()
            f = OxmlElement("w:rFonts")
            f.set(qn("w:cs"), "Arial"); f.set(qn("w:ascii"), "Arial"); f.set(qn("w:hAnsi"), "Arial")
            rf2.append(f)
        return _rtl_par(h)

    def table(header, rows, widths=None):
        t = doc.add_table(rows=1, cols=len(header))
        t.style = "Light Grid Accent 1"
        tblPr = t._tbl.tblPr
        tblPr.append(OxmlElement("w:bidiVisual"))
        for i, h in enumerate(header):
            cell = t.rows[0].cells[i]
            cell.text = ""
            r = cell.paragraphs[0].add_run(h)
            r.bold = True
            r.font.size = Pt(9)
            _rtl_par(cell.paragraphs[0])
        for row in rows:
            cells = t.add_row().cells
            for i, v in enumerate(row):
                cells[i].text = ""
                r = cells[i].paragraphs[0].add_run(str(v))
                r.font.size = Pt(9)
                _rtl_par(cells[i].paragraphs[0])
        if widths:
            for row in t.rows:
                for i, wd in enumerate(widths):
                    row.cells[i].width = Cm(wd)
        return t

    meta, prof = program["meta"], program["profile"]
    head("البرنامج التأهيلي الفردي لعضلات الرقبة", 0)
    para(f"اللاعب: {meta.get('name') or '—'}    الرمز: {meta.get('code') or '—'}    الرياضة: {meta.get('sport') or '—'}", True, 11)
    para(f"تاريخ الإنشاء: {meta.get('created')}    معامل التقدم: {meta.get('progression_factor')}    "
         "8 أسابيع × 3 وحدات (24 وحدة)")
    para(DISCLAIMER)
    head("ملف اللاعب ونتائج القواعد", 1)
    table(["الاختبار/القاعدة", "الدرجة", "التفسير والإجراء"],
          [[f["title"], _SEV_AR[f["severity"]], f["text"]] for f in prof["findings"]], [4, 2, 11.5])
    head("الاحتياطات", 1)
    for p in prof["precautions"]:
        _rtl_par(doc.add_paragraph("• " + p))
    head("الخطة الأسبوعية", 1)
    table(["الأسبوع", "المستوى", "دقائق الوحدات (1/2/3)", "مجموعات × تكرارات", "راحة (ث)", "سقف المقاومة (% MVC)"],
          [[wk["week"], wk["level"], " / ".join(str(x) for x in wk["unit_minutes"]), f"{wk['dose']['sets']} × {wk['dose']['reps']}",
            wk["dose"]["rest_s"], wk["dose"]["resistance_cap_pct"]] for wk in program["weeks"]], [2, 2, 4, 3.5, 2, 3.5])
    used = {}
    for wk in program["weeks"]:
        doc.add_page_break()
        head(f"الأسبوع {wk['week']} — المستوى {wk['level']}", 1)
        for u in wk["units"]:
            head(f"الوحدة {u['unit']}: {u['name']} ({u['minutes']} دقيقة)", 2)
            rows = []
            for blk in ("warmup", "main", "cooldown"):
                for e in u[blk]:
                    used.setdefault(e["id"], e)
                    rows.append([_BLK_AR[blk], e["name"], e["side"] or "—", f"{e['sets']}×{e['reps']}",
                                 e["hold_s"] or "—", e["rest_s"], e["resistance_pct_mvc"] or "—", e["minutes"]])
            table(["الجزء", "التمرين", "الجانب", "مجموعات×تكرارات", "ثبات (ث)", "راحة (ث)", "مقاومة %MVC", "دقائق"],
                  rows, [1.8, 6.0, 1.5, 2.2, 1.4, 1.4, 1.6, 1.3])
    doc.add_page_break()
    head("ملحق: شرح التمارين المستخدمة", 1)
    for ex_id in sorted(used):
        e = used[ex_id]
        para(f"{e['name']}  [{e['category']}]  — المعدات: {e['equipment']}", True)
        for i, s in enumerate(e["instructions"]):
            _rtl_par(doc.add_paragraph(f"{i+1}. {s}"))
        _rtl_par(doc.add_paragraph("أخطاء شائعة: " + "؛ ".join(e["common_errors"])))
    bio = io.BytesIO()
    doc.save(bio)
    return bio.getvalue()

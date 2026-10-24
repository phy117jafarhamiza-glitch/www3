# -*- coding: utf-8 -*-
"""تقرير Word لكل لاعب: القياسات T0–T3، تحليلات EMG، ومجاميع الذكاء الاصطناعي."""
import os
from datetime import datetime

import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement

from utils.constants import STAGES, TEST_NAMES

REPORTS_DIR = "data/reports"


def _rtl(p):
    p._p.get_or_add_pPr().append(OxmlElement("w:bidi"))
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT


def _table(doc, header, rows):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Light Grid Accent 1"
    for i, h in enumerate(header):
        t.rows[0].cells[i].text = str(h)
    for r in rows:
        cells = t.add_row().cells
        for i, v in enumerate(r):
            cells[i].text = "" if v is None or (isinstance(v, float) and pd.isna(v)) else (f"{v:.2f}" if isinstance(v, float) else str(v))


def generate_subject_report(subject, measurements, emg_analyses, ai_sets, out_path=None, sessions=None, coach=None, pain=None):
    os.makedirs(REPORTS_DIR, exist_ok=True)
    doc = Document()
    h = doc.add_heading(f"تقرير اللاعب: {subject['name']}", 1)
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _rtl(h)
    p = doc.add_paragraph(f"الرمز: {subject.get('code') or '-'} | المجموعة: {subject['group_name']} | العمر: {subject.get('age') or '-'} | "
                          f"الرياضة: {subject.get('sport') or '-'} | التاريخ: {datetime.now():%Y-%m-%d}")
    _rtl(p)

    _rtl(doc.add_heading("١. القياسات عبر نقاط الزمن", 2))
    if measurements:
        df = pd.DataFrame(measurements)
        df["المتغير"] = df["test_key"].map(TEST_NAMES)
        piv = df.pivot_table(index="المتغير", columns="stage", values="value", aggfunc="first")
        cols = [s for s in STAGES if s in piv.columns]
        _table(doc, ["المتغير"] + cols, [[i] + [piv.loc[i, c] for c in cols] for i in piv.index])
    else:
        _rtl(doc.add_paragraph("لا توجد قياسات."))

    _rtl(doc.add_heading("٢. تحليلات EMG-Net", 2))
    if emg_analyses:
        _table(doc, ["النقطة", "العضلة", "التصنيف", "الثقة"],
               [[e["stage"], e["muscle"], e["predicted_label"], e["confidence"]] for e in emg_analyses])
    else:
        _rtl(doc.add_paragraph("لا توجد تحليلات."))

    _rtl(doc.add_heading("٣. مجاميع التأهيل الذكي", 2))
    if ai_sets:
        _table(doc, ["الأسبوع", "الجلسة", "الاتجاه", "الجودة", "المقترح %", "النهائي %", "الأمان"],
               [[a["week"], a["session_no"], a["direction"], a["quality_label"], a["rec_resistance"],
                 a["final_resistance"], a["safety_action"]] for a in reversed(ai_sets)])
    else:
        _rtl(doc.add_paragraph("لا توجد مجاميع مسجّلة."))

    _rtl(doc.add_heading("٤. الالتزام بالبرنامج", 2))
    if sessions:
        done = sum(1 for x in sessions if x.get("completed"))
        _rtl(doc.add_paragraph(f"الجلسات المكتملة: {done} من 24 ({100 * done / 24:.0f}%)."))
        _table(doc, ["الأسبوع", "الجلسة", "الدقائق", "RPE", "اكتملت"],
               [[x["week"], x["session_no"], x["minutes"], x["rpe"], "نعم" if x["completed"] else "لا"] for x in sessions])
    else:
        _rtl(doc.add_paragraph("لا توجد جلسات مسجّلة."))

    _rtl(doc.add_heading("٥. المدرّب الحركي (الرؤية الحاسوبية)", 2))
    if coach:
        _table(doc, ["الأسبوع", "التمرين", "التكرارات", "متوسط المدى", "سريعة", "قصيرة", "تعويض", "الجودة"],
               [[c["week"], c["exercise_id"], c["reps"], c["mean_rom"], c["too_fast"], c["short_rom"],
                 c["compensation"], c["quality_score"]] for c in reversed(coach)])
    else:
        _rtl(doc.add_paragraph("لا توجد جلسات مدرّب محفوظة."))

    if pain:
        _rtl(doc.add_heading("٦. الألم اللحظي", 2))
        mx = max(p["nprs"] for p in pain)
        _rtl(doc.add_paragraph(f"عدد التسجيلات: {len(pain)} — أعلى ألم مسجَّل: {mx:g}/10."))

    note = doc.add_paragraph("هذا التقرير أداة مساندة للبحث ولا يغني عن التقييم السريري.")
    _rtl(note)
    out_path = out_path or os.path.join(REPORTS_DIR, f"report_{subject['id']}.docx")
    doc.save(out_path)
    return out_path
